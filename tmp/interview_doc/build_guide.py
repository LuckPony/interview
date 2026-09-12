from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(__file__).with_name('guide.md')
OUT = ROOT / 'output' / 'documents' / '程多多_Java与Agent面试问答_马素超.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(.65)
sec.bottom_margin = Inches(.65)
sec.left_margin = Inches(.77)
sec.right_margin = Inches(.77)
sec.footer_distance = Inches(.28)

def font(style, name, size, bold=False):
    s = doc.styles[style]
    s.font.name = name
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.italic = False
    s.font.color.rgb = RGBColor(0, 0, 0)
    pr = s.element.get_or_add_rPr()
    fonts = pr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); pr.append(fonts)
    for key in ['ascii','hAnsi','eastAsia','cs']:
        fonts.set(qn('w:'+key), name)
    for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        fonts.attrib.pop(qn('w:'+key),None)
    pf=s.paragraph_format
    pf.space_before=Pt(0);pf.space_after=Pt(6)
    pf.line_spacing=1.17
    pf.widow_control=True
    return s

font('Normal','宋体',11)
font('Title','微软雅黑',24,True)
font('Subtitle','微软雅黑',14)
font('Heading 1','微软雅黑',17,True)
font('Heading 2','微软雅黑',12,True)
font('Heading 3','微软雅黑',11.5,True)
font('Caption','宋体',9)
font('Footer','宋体',9)
for st in ['Heading 1','Heading 2','Heading 3']:
    pf=doc.styles[st].paragraph_format
    pf.keep_with_next=True
    pf.space_before=Pt(9 if st!='Heading 1' else 0)
    pf.space_after=Pt(6 if st!='Heading 1' else 14)

def set_run(r,name=None,size=None,bold=None):
    if name:
        r.font.name=name
        r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
    if size:r.font.size=Pt(size)
    if bold is not None:r.bold=bold

def add_text(p,text):
    # Plain labels remain concise; source URLs are clickable in Word.
    m=re.match(r'^(考察点|参考回答|追问|证据准备|答案|题目)：',text)
    if m:
        label=m.group(0)
        rr=p.add_run(label);rr.bold=True
        text=text[len(label):]
    pos=0
    for match in re.finditer(r'https?://[^\s；。]+',text):
        p.add_run(text[pos:match.start()])
        uri=match.group()
        rel=p.part.relate_to(uri,RT.HYPERLINK,is_external=True)
        h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
        run=OxmlElement('w:r');pr=OxmlElement('w:rPr')
        col=OxmlElement('w:color');col.set(qn('w:val'),'333333');pr.append(col)
        run.append(pr);t=OxmlElement('w:t');t.text=uri;run.append(t);h.append(run);p._p.append(h)
        pos=match.end()
    p.add_run(text[pos:])

page=-1
for line in SOURCE.read_text(encoding='utf-8').splitlines():
    if not line:continue
    if line.startswith('@@PAGE '):
        page+=1
        p=doc.add_paragraph(line[7:],style='Title' if page==0 else 'Heading 1')
        if page:p.paragraph_format.page_break_before=True
        continue
    if line.startswith('@SUB '):
        p=doc.add_paragraph(line[5:],style='Subtitle')
        p.paragraph_format.space_after=Pt(16)
    elif line.startswith('@Q '):
        p=doc.add_paragraph(line[3:],style='Heading 2')
    elif line.startswith('@H '):
        p=doc.add_paragraph(line[3:],style='Heading 3')
    elif line.startswith('@SMALL '):
        p=doc.add_paragraph(style='Caption');add_text(p,line[7:])
        p.paragraph_format.space_before=Pt(7)
    elif line.startswith('@CODE '):
        p=doc.add_paragraph()
        p.paragraph_format.left_indent=Inches(.12)
        p.paragraph_format.space_after=Pt(1)
        p.paragraph_format.line_spacing=1.05
        set_run(p.add_run(line[6:]),'Consolas',9.5)
    else:
        p=doc.add_paragraph();add_text(p,line)
        if line.startswith('考察点：'):
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_after=Pt(4)
            for r in p.runs:r.font.size=Pt(9.5)
        if line.startswith('追问：'):
            for r in p.runs:r.font.size=Pt(10.5)
        if page==26:
            for r in p.runs:r.font.size=Pt(9.5)
            p.paragraph_format.line_spacing=1.1
            p.paragraph_format.space_after=Pt(7)

f=sec.footer.paragraphs[0]
f.alignment=WD_ALIGN_PARAGRAPH.CENTER
f.add_run('第 ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
f.add_run(' 页')
doc.core_properties.title='程多多 Java 与 Agent 应用工程师面试问答'
doc.core_properties.subject='马素超面试备考'
doc.core_properties.author='马素超'
doc.core_properties.keywords='程多多 Java Agent homegrown 面试'
for root in [doc.styles.element, doc._element]:
    for border in list(root.xpath('.//w:pBdr')):
        border.getparent().remove(border)
doc.save(OUT)
print(OUT)
print('Questions',len(re.findall(r'^@Q ',SOURCE.read_text(encoding='utf-8'),re.M)))
print('Planned pages',page+1)
