from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path(__file__).with_name("python_guide.md")
OUT = ROOT / "output" / "documents" / "夯智交互_Python后端实习生面试问答_马素超.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.77)
section.right_margin = Inches(0.77)
section.footer_distance = Inches(0.28)


def configure_style(style_name, font_name, size, bold=False):
    style = doc.styles[style_name]
    style.font.name = font_name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), font_name)
    for key in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + key), None)
    fmt = style.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(6)
    fmt.line_spacing = 1.17
    fmt.widow_control = True
    return style


configure_style("Normal", "宋体", 11)
configure_style("Title", "微软雅黑", 24, True)
configure_style("Subtitle", "微软雅黑", 14)
configure_style("Heading 1", "微软雅黑", 17, True)
configure_style("Heading 2", "微软雅黑", 12, True)
configure_style("Heading 3", "微软雅黑", 11.5, True)
configure_style("Caption", "宋体", 9)
configure_style("Footer", "宋体", 9)

for style_name in ("Heading 1", "Heading 2", "Heading 3"):
    fmt = doc.styles[style_name].paragraph_format
    fmt.keep_with_next = True
    fmt.space_before = Pt(9 if style_name != "Heading 1" else 0)
    fmt.space_after = Pt(6 if style_name != "Heading 1" else 14)


def set_run_font(run, name=None, size=None, bold=None):
    if name:
        run.font.name = name
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_hyperlinked_text(paragraph, text):
    match = re.match(r"^(考察点|参考回答|追问|证据准备|回答边界|答案|题目)：", text)
    if match:
        label = match.group(0)
        label_run = paragraph.add_run(label)
        label_run.bold = True
        text = text[len(label):]
    cursor = 0
    for url_match in re.finditer(r"https?://[^\s；。]+", text):
        paragraph.add_run(text[cursor:url_match.start()])
        url = url_match.group()
        relationship = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), relationship)
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "333333")
        properties.append(color)
        run.append(properties)
        text_element = OxmlElement("w:t")
        text_element.text = url
        run.append(text_element)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)
        cursor = url_match.end()
    paragraph.add_run(text[cursor:])


page_index = -1
source_text = SOURCE.read_text(encoding="utf-8")
for line in source_text.splitlines():
    if not line:
        continue
    if line.startswith("@@PAGE "):
        page_index += 1
        paragraph = doc.add_paragraph(
            line[7:], style="Title" if page_index == 0 else "Heading 1"
        )
        if page_index:
            paragraph.paragraph_format.page_break_before = True
        continue
    if line.startswith("@SUB "):
        paragraph = doc.add_paragraph(line[5:], style="Subtitle")
        paragraph.paragraph_format.space_after = Pt(16)
    elif line.startswith("@Q "):
        doc.add_paragraph(line[3:], style="Heading 2")
    elif line.startswith("@H "):
        doc.add_paragraph(line[3:], style="Heading 3")
    elif line.startswith("@SMALL "):
        paragraph = doc.add_paragraph(style="Caption")
        add_hyperlinked_text(paragraph, line[7:])
        paragraph.paragraph_format.space_before = Pt(7)
    elif line.startswith("@CODE "):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.12)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.paragraph_format.line_spacing = 1.05
        set_run_font(paragraph.add_run(line[6:]), "Consolas", 9.3)
    else:
        paragraph = doc.add_paragraph()
        add_hyperlinked_text(paragraph, line)
        if line.startswith("考察点："):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(4)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
        if line.startswith("追问："):
            for run in paragraph.runs:
                run.font.size = Pt(10.5)
        if page_index == 27:
            for run in paragraph.runs:
                run.font.size = Pt(10)
            paragraph.paragraph_format.line_spacing = 1.08
            paragraph.paragraph_format.space_after = Pt(4)
        if page_index == 28:
            for run in paragraph.runs:
                run.font.size = Pt(9)
            paragraph.paragraph_format.line_spacing = 1.03
            paragraph.paragraph_format.space_after = Pt(3)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("第 ")
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
footer.add_run(" 页")

doc.core_properties.title = "上海夯智交互文化科技 Python 后端实习生面试问答"
doc.core_properties.subject = "马素超面试备考"
doc.core_properties.author = "马素超"
doc.core_properties.keywords = "Python FastAPI Agent Coding Harness 面试"

# Remove inherited Word title borders and theme decorations.
for root in (doc.styles.element, doc._element):
    for border in list(root.xpath(".//w:pBdr")):
        border.getparent().remove(border)

doc.save(OUT)
print(OUT)
print("Questions", len(re.findall(r"^@Q ", source_text, re.M)))
print("Planned pages", page_index + 1)
